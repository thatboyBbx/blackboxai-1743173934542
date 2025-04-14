import PyPDF2
import docx
from openpyxl import load_workbook
from pptx import Presentation
from .error_handlers import TextExtractionError

def extract_text_from_xlsx(file_path):
    """Extract text from Excel files"""
    try:
        text = []
        wb = load_workbook(file_path)
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            for row in ws.iter_rows(values_only=True):
                text.append(" ".join(str(cell) for cell in row if cell))
        return "\n".join(text)
    except Exception as e:
        raise TextExtractionError(file_path, f"Excel extraction failed: {str(e)}")

def extract_text_from_pptx(file_path):
    """Extract text from PowerPoint files"""
    try:
        text = []
        prs = Presentation(file_path)
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text.append(shape.text)
        return "\n".join(text)
    except Exception as e:
        raise TextExtractionError(file_path, f"PowerPoint extraction failed: {str(e)}")

def extract_text_from_pdf(file_path):
    """Extract text from PDF files"""
    try:
        text = ""
        with open(file_path, 'rb') as file:
            reader = PyPDF2.PdfReader(file)
            for page in reader.pages:
                text += page.extract_text() + "\n"
        return text
    except Exception as e:
        raise TextExtractionError(file_path, f"PDF extraction failed: {str(e)}")

def extract_text_from_docx(file_path):
    """Extract text from DOCX files"""
    try:
        doc = docx.Document(file_path)
        return "\n".join([para.text for para in doc.paragraphs])
    except Exception as e:
        raise TextExtractionError(file_path, f"DOCX extraction failed: {str(e)}")

def get_document_text(file_path):
    """Main text extraction interface"""
    if file_path.endswith('.pdf'):
        return extract_text_from_pdf(file_path)
    elif file_path.endswith('.docx'):
        return extract_text_from_docx(file_path)
    elif file_path.endswith('.xlsx'):
        return extract_text_from_xlsx(file_path)
    elif file_path.endswith('.pptx'):
        return extract_text_from_pptx(file_path)
    elif file_path.endswith('.txt'):
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()
    else:
        raise TextExtractionError(file_path, "Unsupported file format")
